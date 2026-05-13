"""Row iterators for each upstream source the OBGYN classifier filters.

Each iterator yields ``(row_id, user_message_text)`` tuples in source order.
The ``user_message_text`` is the exact string the classifier driver will send
as the user message — per-source rendering rules are encoded here, not in the
driver.

Source identifiers are taken from the upstream rows where present
(HealthBench's ``prompt_id``, Kenya's ``StudyID``); MedQA-USMLE has no native
row id so we synthesise one from the row index.
"""

from __future__ import annotations

import json
from pathlib import Path
from typing import Iterator, Sequence


def render_healthbench_messages(messages: Sequence[dict]) -> str:
    """Render a HealthBench prompt (list of role/content messages) as a single
    user-message string with ``User:`` / ``Assistant:`` role prefixes.

    HealthBench rows always start and end with a ``user`` message and strictly
    alternate roles in between, but we do not assume that here — any role
    sequence is rendered verbatim with the role title-cased.
    """
    parts = []
    for m in messages:
        role = str(m["role"]).capitalize()
        parts.append(f"{role}: {m['content']}")
    return "\n".join(parts)


def render_medqa_options(question: str, options: dict[str, str]) -> str:
    """Render a MedQA-USMLE question + its answer options as a single string.

    Options are sorted by letter and joined with ` | ` after a literal
    ``Options: `` prefix on a new line — matching the contract the MCQ-mode
    classifier prompt expects.
    """
    opts_str = " | ".join(f"{letter}. {text}" for letter, text in sorted(options.items()))
    return f"{question}\nOptions: {opts_str}"


def iter_healthbench(jsonl_path: Path) -> Iterator[tuple[str, str]]:
    """Yield ``(prompt_id, rendered_conversation)`` for each HealthBench row."""
    with Path(jsonl_path).open(encoding="utf-8") as f:
        for line in f:
            line = line.strip()
            if not line:
                continue
            row = json.loads(line)
            yield row["prompt_id"], render_healthbench_messages(row["prompt"])


def iter_medqa_usmle(jsonl_path: Path) -> Iterator[tuple[str, str]]:
    """Yield ``(row_id, question + options)`` for each MedQA-USMLE row.

    MedQA-USMLE source rows have no native id; we synthesise ``usmle_{i:05d}``
    using the zero-padded row index for stable, sortable identifiers.
    """
    with Path(jsonl_path).open(encoding="utf-8") as f:
        for i, line in enumerate(f):
            line = line.strip()
            if not line:
                continue
            row = json.loads(line)
            text = render_medqa_options(row["question"], row.get("options", {}))
            yield f"usmle_{i:05d}", text


def iter_kenya(xlsx_path: Path) -> Iterator[tuple[str, str]]:
    """Yield ``(study_id, scenario)`` for each Kenya source vignette.

    Reads the `Prompt responses.xlsx` artifact directly with openpyxl. The
    ``StudyID`` column is the stable upstream identifier; ``User Prompt`` is
    the nurse-written clinical scenario rendered as the user message.
    """
    try:
        from openpyxl import load_workbook
    except ImportError as e:
        raise ImportError(
            "openpyxl is required to read the Kenya source xlsx — "
            "install with `pip install openpyxl`"
        ) from e

    wb = load_workbook(filename=str(xlsx_path), read_only=True, data_only=True)
    sheet = wb.active
    rows = sheet.iter_rows(values_only=True)
    header = next(rows)
    col_index = {name: idx for idx, name in enumerate(header)}
    for required in ("StudyID", "User Prompt"):
        if required not in col_index:
            raise ValueError(
                f"missing required column {required!r} in {xlsx_path}; "
                f"got columns {list(header)}"
            )
    sid_idx = col_index["StudyID"]
    prompt_idx = col_index["User Prompt"]
    for row in rows:
        sid = row[sid_idx]
        prompt = row[prompt_idx]
        if sid is None or prompt is None:
            continue
        yield str(sid), str(prompt)

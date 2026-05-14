"""Adapter for Kenya Clinical Vignettes source xlsx.

Reads the upstream `Prompt responses.xlsx` (all 507 source vignettes —
not the pre-filtered TSV), filters by the OBGYN classifier verdicts
produced upstream, and emits v0.4 `open_ended` rows where ``question`` is
the nurse-written scenario and ``answer`` is the Kenyan clinician's
reference response.

openpyxl is imported lazily so the module is importable without it; only
``iter_kenya_source`` actually requires it.
"""

from __future__ import annotations

import json
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Iterator, Mapping


KENYA_SOURCE_DATASET = "Kenya-Clinical-Vignettes"
KENYA_SOURCE_URL = (
    "https://www.medrxiv.org/content/10.1101/2025.10.25.25338798v1"
)
KENYA_LICENSE = "MIT"

SCHEMA_VERSION = "0.4"
VALID_CATEGORIES = frozenset(
    {"MATERNAL", "NEONATAL", "CHILD_HEALTH", "SEXUAL_AND_REPRODUCTIVE_HEALTH", "NONE"}
)


class KenyaAdapterError(ValueError):
    """Raised when a Kenya row cannot be normalized."""


@dataclass
class KenyaSourceRow:
    """One row read from the Kenya source xlsx."""

    study_id: str
    scenario: str
    clinician_response: str


@dataclass
class KenyaAdapterStats:
    """Filter accounting; invariant: total == kept + skipped_none + skipped_no_verdict."""

    total: int = 0
    kept: int = 0
    skipped_none: int = 0
    skipped_no_verdict: int = 0

    def as_dict(self) -> dict[str, int]:
        return {
            "total": self.total,
            "kept": self.kept,
            "skipped_none": self.skipped_none,
            "skipped_no_verdict": self.skipped_no_verdict,
        }


def iter_kenya_source(xlsx_path: str | Path) -> Iterator[KenyaSourceRow]:
    """Yield one ``KenyaSourceRow`` per vignette in `Prompt responses.xlsx`.

    Requires openpyxl; raises ImportError with an install hint if it isn't
    available. Rows with a null StudyID, scenario, or clinician_response
    are skipped silently (trailing empty rows are common in xlsx exports).
    """
    try:
        from openpyxl import load_workbook
    except ImportError as e:
        raise ImportError(
            "openpyxl is required to read the Kenya source xlsx — "
            "install with `pip install openpyxl`"
        ) from e

    workbook = load_workbook(filename=str(xlsx_path), read_only=True, data_only=True)
    sheet = workbook.active
    rows_iter = sheet.iter_rows(values_only=True)
    header = next(rows_iter)
    col_index = {name: idx for idx, name in enumerate(header)}
    required_columns = ("StudyID", "User Prompt", "Clinician response")
    for col in required_columns:
        if col not in col_index:
            raise KenyaAdapterError(
                f"missing required column {col!r} in {xlsx_path}; got columns {list(header)}"
            )
    sid_idx = col_index["StudyID"]
    prompt_idx = col_index["User Prompt"]
    response_idx = col_index["Clinician response"]
    for row in rows_iter:
        sid = row[sid_idx]
        prompt = row[prompt_idx]
        response = row[response_idx]
        if sid is None or prompt is None or response is None:
            continue
        yield KenyaSourceRow(
            study_id=str(sid),
            scenario=str(prompt),
            clinician_response=str(response),
        )


def load_verdicts(path: str | Path) -> dict[str, dict[str, Any]]:
    """Read the Kenya OBGYN classifier verdicts JSONL into ``{study_id: verdict}``."""
    verdicts: dict[str, dict[str, Any]] = {}
    with Path(path).open(encoding="utf-8") as handle:
        for line in handle:
            line = line.strip()
            if not line:
                continue
            verdict = json.loads(line)
            row_id = verdict.get("row_id")
            category = verdict.get("category")
            if not isinstance(row_id, str) or category not in VALID_CATEGORIES:
                continue
            verdicts[row_id] = verdict
    return verdicts


def normalize_kenya_row(
    source_row: KenyaSourceRow,
    *,
    verdict: Mapping[str, Any],
    benchmark_version: str,
    key_facts_metadata: Mapping[str, Any] | None = None,
) -> dict[str, Any]:
    """Build one v0.4 ``open_ended`` mamabench row from a Kenya source row.

    When ``key_facts_metadata`` is provided (loaded from the keyfact extractor
    side-file by row_id), it is nested under
    ``source.metadata.key_fact_extraction`` alongside ``obgyn_classification``.
    """
    scenario = source_row.scenario.strip()
    response = source_row.clinician_response.strip()
    if not scenario:
        raise KenyaAdapterError(f"StudyID {source_row.study_id}: empty scenario")
    if not response:
        raise KenyaAdapterError(f"StudyID {source_row.study_id}: empty clinician_response")
    metadata: dict[str, Any] = {
        "obgyn_classification": {
            "model": verdict.get("model"),
            "prompt_version": verdict.get("prompt_version"),
            "category": verdict["category"],
            "rationale": verdict.get("rationale", ""),
        },
    }
    if key_facts_metadata is not None:
        metadata["key_fact_extraction"] = dict(key_facts_metadata)
    return {
        "id": f"mamabench_{benchmark_version}_kenya_{source_row.study_id}",
        "schema_version": SCHEMA_VERSION,
        "set_type": "open_ended",
        "question": scenario,
        "answer": response,
        "source": {
            "dataset": KENYA_SOURCE_DATASET,
            "id": source_row.study_id,
            "metadata": metadata,
        },
    }


def load_kenya(
    source_xlsx: str | Path,
    verdicts_path: str | Path,
    *,
    benchmark_version: str,
    limit: int | None = None,
    keyfacts_path: str | Path | None = None,
) -> tuple[list[dict[str, Any]], KenyaAdapterStats]:
    """Walk the Kenya source xlsx, filter by verdict, normalize to v0.4 rows.

    When ``keyfacts_path`` is provided, the keyfact extractor side-file is
    loaded and each row matched by id gets
    ``source.metadata.key_fact_extraction`` populated.
    """
    if limit is not None and limit < 0:
        raise KenyaAdapterError("limit must be non-negative")

    keyfacts_by_id: dict[str, dict[str, Any]] = {}
    if keyfacts_path is not None:
        from ._keyfact_extraction import load_keyfacts_by_row_id

        keyfacts_by_id = load_keyfacts_by_row_id(keyfacts_path)

    verdicts = load_verdicts(verdicts_path)
    stats = KenyaAdapterStats()
    rows: list[dict[str, Any]] = []
    for source_row in iter_kenya_source(source_xlsx):
        stats.total += 1
        verdict = verdicts.get(source_row.study_id)
        if verdict is None:
            stats.skipped_no_verdict += 1
            continue
        if verdict["category"] == "NONE":
            stats.skipped_none += 1
            continue
        row_id = f"mamabench_{benchmark_version}_kenya_{source_row.study_id}"
        rows.append(
            normalize_kenya_row(
                source_row,
                verdict=verdict,
                benchmark_version=benchmark_version,
                key_facts_metadata=keyfacts_by_id.get(row_id),
            )
        )
        stats.kept += 1
        if limit is not None and len(rows) >= limit:
            break

    if stats.total != stats.kept + stats.skipped_none + stats.skipped_no_verdict:
        raise KenyaAdapterError(
            f"filter accounting inconsistent — total={stats.total}, "
            f"kept={stats.kept}, skipped_none={stats.skipped_none}, "
            f"skipped_no_verdict={stats.skipped_no_verdict}"
        )
    return rows, stats


def build_kenya_source_metadata() -> dict[str, Any]:
    """Build dataset-level Kenya provenance for the manifest."""
    return {
        KENYA_SOURCE_DATASET: {
            "url": KENYA_SOURCE_URL,
            "license": KENYA_LICENSE,
        }
    }
